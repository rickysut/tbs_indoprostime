import os
import re

import frappe
from openpyxl import load_workbook

DEFAULT_COUNTRY = "ID"
ADDRESS_TYPE = "Office"


def run(file_path=None, dry_run=False, start_from_code=None):
    """Migrate rows from VENDOR-asli.xlsx to Business Partner with Address & Contact.

    Run with:
        bench --site <site> execute tbs_indoprostime.scripts.migrate_vendor_to_business_partner.run
    Or with dry-run:
        bench --site <site> execute tbs_indoprostime.scripts.migrate_vendor_to_business_partner.run --kwargs '{"dry_run": True}'
    Resume from a specific Code:
        bench --site <site> execute tbs_indoprostime.scripts.migrate_vendor_to_business_partner.run --kwargs '{"start_from_code": "0005"}'
    """
    start_from_code = str(start_from_code) if start_from_code else None
    if not file_path:
        file_path = _default_file_path()

    wb = load_workbook(file_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h else "" for h in rows[0]]
    data_rows = rows[1:]

    col_map = {name: idx for idx, name in enumerate(header)}

    total = 0
    created = 0
    updated = 0
    failed = 0

    missing_industries = set()
    missing_payment_terms = set()

    for row in data_rows:
        total += 1
        code = _str(row, col_map, "Code")
        bp_code = f"V{code}" if code else ""
        name = _str(row, col_map, "Name") or bp_code

        print(f"[{total}] Processing {code} -> {bp_code} - {name}")

        if not code:
            failed += 1
            msg = f"Row {total}: Code is empty"
            frappe.log_error(title="Migrate Vendor to BP - Empty Code", message=msg)
            print(f"  -> FAILED: {msg}")
            continue

        if start_from_code and code < start_from_code:
            print(f"  -> SKIP (before {start_from_code})")
            continue

        bp = None
        try:
            if frappe.db.exists("Business Partner", bp_code):
                bp = frappe.get_doc("Business Partner", bp_code)
                _apply_business_partner_values(bp, row, col_map, missing_industries, missing_payment_terms)
                if dry_run:
                    print(f"  -> DRY RUN: would update BP {bp_code}")
                    continue
                bp.save(ignore_permissions=True)
                updated += 1
                print(f"  -> UPDATED: {bp.name}")
            else:
                bp = _create_business_partner(
                    row, col_map, missing_industries, missing_payment_terms, dry_run
                )
                if dry_run:
                    print(f"  -> DRY RUN: would create BP {bp_code}")
                    continue
                created += 1
                print(f"  -> CREATED: {bp.name}")

            if not bp.primary_address:
                try:
                    _create_address(row, col_map, bp)
                except Exception as e:
                    frappe.log_error(
                        title="Migrate Vendor to BP - Address Error",
                        message=f"{bp_code}: {e}",
                    )
                    print(f"  -> ADDRESS ERROR: {e}")

            if not bp.primary_contact:
                try:
                    _create_contact(row, col_map, bp)
                except Exception as e:
                    frappe.log_error(
                        title="Migrate Vendor to BP - Contact Error",
                        message=f"{bp_code}: {e}",
                    )
                    print(f"  -> CONTACT ERROR: {e}")
                    _append_phone_to_description(bp, row, col_map)

            if bp.primary_address or bp.primary_contact:
                try:
                    bp.save(ignore_permissions=True)
                except Exception as e:
                    frappe.log_error(
                        title="Migrate Vendor to BP - Save BP Error",
                        message=f"{bp_code}: {e}",
                    )
                    print(f"  -> SAVE BP ERROR: {e}")

            frappe.db.commit()

        except Exception as e:
            failed += 1
            frappe.log_error(
                title="Migrate Vendor to BP - BP Creation Error",
                message=f"{bp_code}: {e}",
            )
            print(f"  -> FAILED: {e}")

    if missing_industries:
        frappe.log_error(
            title="Migrate Vendor to BP - Missing Industries",
            message="Industries not found:\n" + "\n".join(missing_industries),
        )
    if missing_payment_terms:
        frappe.log_error(
            title="Migrate Vendor to BP - Missing Payment Terms",
            message="Payment Terms not found:\n" + "\n".join(missing_payment_terms),
        )

    summary = (
        f"Migrate Vendor to Business Partner completed: "
        f"total={total}, created={created}, updated={updated}, failed={failed}"
    )
    print(summary)
    if not dry_run:
        frappe.log_error(title="Migrate Vendor to Business Partner - Summary", message=summary)


def _create_business_partner(row, col_map, missing_industries, missing_payment_terms, dry_run):
    if dry_run:
        return None

    bp = frappe.new_doc("Business Partner")
    _apply_business_partner_values(bp, row, col_map, missing_industries, missing_payment_terms)
    bp.insert(ignore_permissions=True)
    return bp


def _apply_business_partner_values(bp, row, col_map, missing_industries, missing_payment_terms):
    code = _str(row, col_map, "Code")
    bp_code = f"V{code}" if code else ""

    tax_id = _clean_npwp(_str(row, col_map, "NPWP"))
    bp_description = _str(row, col_map, "Comment\n")

    industry = _str(row, col_map, "Industry")
    if industry and not frappe.db.exists("Industry Type", industry):
        missing_industries.add(industry)
        industry = None

    payment_term = _str(row, col_map, "TERMOFPAY,N,3,0")
    if payment_term and not frappe.db.exists("Payment Terms Template", payment_term):
        missing_payment_terms.add(payment_term)
        payment_term = None

    bp.update(
        {
            "bp_code": bp_code,
            "partner_name": _str(row, col_map, "Name") or bp_code,
            "partner_type": "Company",
            "vendor": 1,
            "customer": 0,
            "vendor_group": "Forwarding",
            "tax_id": tax_id,
            "bp_description": bp_description,
            "industry": industry,
            "buying_payment_term": payment_term,
        }
    )


def _create_address(row, col_map, bp):
    raw_lines = []
    for col in ["Address 1", "Address 2", "Address 3", "Address 4"]:
        val = _str(row, col_map, col)
        if val:
            raw_lines.append(val)

    city = _str(row, col_map, "City")
    pincode = _str(row, col_map, "Postalcode")
    country = _str(row, col_map, "Country")
    if not country or not frappe.db.exists("Country", country):
        country = DEFAULT_COUNTRY

    if not city:
        city = "N/A"

    address_line1 = raw_lines[0] if raw_lines else bp.partner_name
    address_line2 = ", ".join(raw_lines[1:]) if len(raw_lines) > 1 else None

    address = frappe.new_doc("Address")
    address.update(
        {
            "address_type": ADDRESS_TYPE,
            "address_line1": address_line1,
            "address_line2": address_line2,
            "city": city,
            "pincode": pincode,
            "country": country,
            "links": [
                {
                    "link_doctype": "Business Partner",
                    "link_name": bp.bp_code,
                }
            ],
        }
    )
    address.insert(ignore_permissions=True)
    bp.primary_address = address.name


def _create_contact(row, col_map, bp):
    contact_name = _str(row, col_map, "Contact")
    phone = _str(row, col_map, "Phone")
    fax = _str(row, col_map, "Fax")

    if not contact_name and not phone and not fax:
        return

    first_name, last_name = _split_name(contact_name or bp.partner_name)

    contact = frappe.new_doc("Contact")
    contact.update(
        {
            "first_name": first_name,
            "last_name": last_name,
            "links": [
                {
                    "link_doctype": "Business Partner",
                    "link_name": bp.bp_code,
                }
            ],
        }
    )

    if frappe.db.has_column("Contact", "custom_fax"):
        contact.custom_fax = fax

    if phone:
        contact.append("phone_nos", {"phone": phone, "is_primary_phone": 1})

    contact.insert(ignore_permissions=True)
    bp.primary_contact = contact.name


def _default_file_path():
    # apps/tbs_indoprostime/tbs_indoprostime/scripts -> apps -> bench root
    bench_root = os.path.dirname(os.path.dirname(os.path.dirname(frappe.get_app_path("tbs_indoprostime"))))
    return os.path.join(bench_root, ".opencode", "img", "VENDOR-asli.xlsx")


def _str(row, col_map, col_name):
    idx = col_map.get(col_name)
    if idx is None:
        return ""
    val = row[idx]
    if val is None:
        return ""
    return str(val).strip()


def _clean_npwp(val):
    if not val:
        return None
    cleaned = re.sub(r"[^0-9]", "", str(val))
    return cleaned if cleaned else None


def _append_phone_to_description(bp, row, col_map):
    phone = _str(row, col_map, "Phone")
    if not phone:
        return
    existing = (bp.bp_description or "").strip()
    new_line = f"Phone: {phone}"
    if new_line in existing:
        return
    bp.bp_description = f"{existing}\n{new_line}" if existing else new_line
    try:
        bp.save(ignore_permissions=True)
    except Exception as e:
        frappe.log_error(
            title="Migrate Vendor to BP - Save Description Error",
            message=f"{bp.bp_code}: {e}",
        )


def _split_name(full_name):
    full_name = (full_name or "").strip()
    if not full_name:
        return "-", ""
    parts = full_name.split(None, 1)
    first = parts[0]
    last = parts[1] if len(parts) > 1 else ""
    return first, last
